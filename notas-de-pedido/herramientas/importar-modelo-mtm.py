#!/usr/bin/env python3
# =============================================================
# importar-modelo-mtm.py
# Convierte la planilla "Modelo MTM Zinguería.xlsx" en datos/tiempos-mtm.json
# (los parámetros, actividades, materiales, tipos de pliegue, detalles y factor de peso
# que usa src/logica/tiempos.js).
#
# Es una herramienta de desarrollo (no la usa el programa en el navegador). Necesita
# Python 3 y openpyxl:  python3 -I herramientas/importar-modelo-mtm.py "Modelo MTM Zinguería.xlsx"
#
# Para cambiar SOLO valores sin Python: en Google Sheets, hoja PARAMETROS_API →
# Archivo > Descargar > CSV, y en el programa: Producción > Parámetros de tiempos > Importar CSV.
#
# Los valores CALCULADOS de la planilla (suplemento total, tiempo por golpe de guillotina,
# desarrollo máximo, ciclo manual, largo máximo) NO se copian: el programa los calcula, para
# que haya una sola fuente. El ancho de chapa y el despunte se toman de datos/config.json.
# =============================================================
import json
import sys
from pathlib import Path

import openpyxl

CALCULADOS = {
    'p_supl_total': 'suma de los suplementos (p_supl_np + … + p_supl_demoras)',
    'p_t_golpe_guill_s': '60 ÷ p_guill_gpm',
    'p_d_max_mm': 'ancho de chapa (config.anchoChapaCm) − despunte (config.cantidadDeChapas.margenEscuadreCm)',
    'p_t_ciclo_manual': 'actividad A21 × p_tmu_s',
    'p_l_max_m': 'el menor entre p_guill_largo_max_m y p_pleg_largo_max_m',
}
DESDE_CONFIG = {
    'p_ancho_chapa_mm': 'config.anchoChapaCm × 10',
    'p_margen_escuadre_mm': 'config.cantidadDeChapas.margenEscuadreCm × 10',
}
# Parámetros que agrega el PROGRAMA para pasar un renglón de la nota a las entradas del modelo
# (la planilla los pide cargados a mano: tipo de perfil, pliegues por tipo, detalles). Todos [SUPUESTO].
CRITERIO = 'Criterio del programa para leer la nota (no está en la planilla). Revisar con el jefe de producción.'
PARAMETROS_PROGRAMA = [
    {'clave': 'q_tolerancia_angulo_recto_grados', 'valor': 1.0, 'unidad': 'grados',
     'descripcion': 'Un pliegue a 90° ± esta tolerancia cuenta como SIMPLE; con otro ángulo, como ANGULO.'},
    {'clave': 'q_tapas_por_tapa', 'valor': 1.0, 'unidad': 'TAPA',
     'descripcion': 'Detalle TAPA que se carga por cada tapa puesta en la canaleta.'},
    {'clave': 'q_soldaduras_por_boqueta', 'valor': 1.0, 'unidad': 'uniones',
     'descripcion': 'Uniones de SOLDADURA por cada boqueta puesta en la canaleta (además del corte del agujero).'},
    {'clave': 'q_encastres_por_corte', 'valor': 1.0, 'unidad': 'ENCASTRE',
     'descripcion': 'Cortes con tijera (ENCASTRE) por cada agujero: el de una boqueta o un corte pedido.'},
    {'clave': 'q_soldaduras_por_pared', 'valor': 1.0, 'unidad': 'uniones',
     'descripcion': 'Uniones de SOLDADURA por cada corte que lleva pared o aletas alrededor.'},
    {'clave': 'q_remaches_por_union', 'valor': 4.0, 'unidad': 'remaches',
     'descripcion': 'Remaches por cada unión de empalme (piezas más largas que la chapa).'},
    {'clave': 'q_sellado_por_union', 'valor': 1.0, 'unidad': '0/1',
     'descripcion': '1 = cada unión de empalme se sella a lo largo de todo el desarrollo; 0 = no se sella.'},
    {'clave': 'q_pendiente_es_especial', 'valor': 1.0, 'unidad': '0/1',
     'descripcion': '1 = un plegado con pendiente se trata como perfil ESPECIAL (topes distintos en cada punta).'},
]
for parametro in PARAMETROS_PROGRAMA:
    parametro.update({'fuente': CRITERIO, 'confianza': 'Baja', 'estado': '[SUPUESTO]', 'grupo': 'Programa'})

# Productos sin perfil plegado (boquetas sueltas, sombreros, codos, artículos de reventa):
# el modelo MTM no los cubre, así que llevan un tiempo propio por grupo del catálogo. [SUPUESTO]
PRODUCTOS_SIN_PERFIL = [
    {'grupo': 'boquetas', 'setupS': 120, 'piezaS': 240, 'maquina': 'banco', 'estado': '[SUPUESTO]',
     'fuente': 'Estimación: desarrollo + engrafe + soldadura de 2 a 3 uniones (p_t_soldadura_union_s = 60 s).'},
    {'grupo': 'canos', 'setupS': 180, 'piezaS': 420, 'maquina': 'banco', 'estado': '[SUPUESTO]',
     'fuente': 'Estimación: codos y ramales armados por gajos con engrafe y soldadura.'},
    {'grupo': 'tapas', 'setupS': 60, 'piezaS': 120, 'maquina': 'banco', 'estado': '[A VALIDAR]',
     'fuente': 'Igual a p_t_tapa_s (planilla MTM).'},
    {'grupo': 'sombreros', 'setupS': 300, 'piezaS': 900, 'maquina': 'banco', 'estado': '[SUPUESTO]',
     'fuente': 'Estimación: varias piezas cilindradas, engrafadas y remachadas.'},
    {'grupo': 'articulos', 'setupS': 0, 'piezaS': 30, 'maquina': 'deposito', 'estado': '[SUPUESTO]',
     'fuente': 'Artículo de reventa: solo buscarlo y prepararlo (no se fabrica).'},
    {'grupo': '*', 'setupS': 180, 'piezaS': 600, 'maquina': 'banco', 'estado': '[SUPUESTO]',
     'fuente': 'Valor por defecto para cualquier otro producto por unidad.'},
]

# Material de la planilla → material y terminación del programa (datos/materiales.json).
FAMILIAS = {
    'GALV': ('galvanizado', 'N'), 'PREP': ('galvanizado', 'P'), 'CINC': ('cincalum', 'N'),
    'ALUM': ('aluminio', 'N'), 'INOX': ('inoxidable', 'N'), 'COBRE': ('cobre', 'N'),
}


def filas(hoja, desde, hasta_col):
    for fila in hoja.iter_rows(min_row=desde, values_only=True):
        if fila[0] is None:
            continue
        yield fila[:hasta_col]


def numero(valor):
    return float(valor) if isinstance(valor, (int, float)) else valor


def main(ruta):
    libro = openpyxl.load_workbook(ruta, data_only=True)

    parametros = []
    for clave, valor, unidad, descripcion, fuente, confianza, estado, grupo in filas(libro['PARAMETROS_GLOBALES'], 5, 8):
        if not str(clave).startswith('p_'):
            continue
        if clave in CALCULADOS or clave in DESDE_CONFIG:
            continue
        parametros.append({
            'clave': clave, 'valor': numero(valor), 'unidad': unidad, 'descripcion': descripcion,
            'fuente': fuente, 'confianza': confianza, 'estado': estado, 'grupo': grupo,
        })

    actividades = []
    for fila in libro['ACTIVIDADES'].iter_rows(min_row=5, values_only=True):
        ident, nombre, clase, nivel, justificacion, np_, mp, nmp = fila[13:21]
        if not ident or not str(ident).startswith('A'):
            continue
        actividades.append({
            'id': ident, 'nombre': nombre, 'clase': clase, 'nivel': nivel, 'justificacion': justificacion,
            'tmuSinPeso': numero(np_), 'tmuConPeso': numero(mp or 0), 'movimientosConPeso': numero(nmp or 0),
        })

    materiales = []
    fuente_anterior = ''
    for fila in filas(libro['MATERIALES'], 5, 21):
        codigo = fila[0]
        if not codigo or '-' not in str(codigo):
            continue
        familia = str(codigo).split('-')[0]
        material_id, terminacion = FAMILIAS[familia]
        materiales.append({
            'codigo': codigo, 'nombre': fila[1], 'materialId': material_id, 'terminacion': terminacion,
            'espesorMm': numero(fila[2]), 'densidadKgM3': numero(fila[4]), 'recubrimientoKgM2': numero(fila[5]),
            'rmMpa': numero(fila[7]), 'factorManipulacion': numero(fila[8]), 'factorPlegado': numero(fila[9]),
            'factorCorte': numero(fila[10]),
            'fuente': fuente_anterior if str(fila[19]).startswith('Idem') else fila[19], 'estado': fila[20],
        })
        if not str(fila[19]).startswith('Idem'):
            fuente_anterior = fila[19]

    tipos_pliegue = []
    for codigo, descripcion, golpes, volteos, factor, herramental, justificacion, estado in filas(libro['TIPOS_PLIEGUE_Y_DETALLES'], 5, 8):
        if codigo in ('Código',) or codigo is None:
            continue
        if codigo.startswith('Vector'):
            break
        tipos_pliegue.append({
            'codigo': codigo, 'descripcion': descripcion, 'golpes': numero(golpes), 'volteos': numero(volteos),
            'factor': numero(factor), 'herramental': numero(herramental), 'justificacion': justificacion, 'estado': estado,
        })

    # Detalles: el tiempo de proceso se guarda como FÓRMULA de parámetros (una sola fuente).
    # proceso = numerador ÷ denominador (número o clave de parámetro; sin denominador = ÷ 1).
    proceso = {
        'REMACHE': {'proceso': {'numerador': 'p_t_perforado_s'}}, 'ENCASTRE': {'proceso': {'numerador': 0}},
        'SELLADO': {'proceso': {'numerador': 1, 'denominador': 'p_v_sellado_m_s'}},
        'TAPA': {'proceso': {'numerador': 'p_t_tapa_s'}},
        'SOLDADURA': {'proceso': {'numerador': 'p_t_soldadura_union_s'}},
        'CURVADO': {'proceso': {'numerador': 'p_pasadas_cilindrado', 'denominador': 'p_v_cilindrado_m_s'}},
        'ENGRAFE': {'proceso': {'numerador': 1, 'denominador': 'p_v_engrafe_m_s'}},
    }
    detalles = []
    for fila in filas(libro['TIPOS_PLIEGUE_Y_DETALLES'], 21, 9):
        codigo = fila[0]
        if codigo not in proceso:
            continue
        detalles.append({
            'codigo': codigo, 'descripcion': fila[1], 'unidad': fila[2],
            'actividadPorUnidad': fila[3] or None, 'actividadFijaPorPieza': fila[5] or None,
            **proceso[codigo], 'estado': fila[8],
        })

    factor_peso = []
    for fila in libro['MTM_CODIGOS'].iter_rows(min_row=5, max_row=16, values_only=True):
        desde, hasta, sc, df = fila[9:13]
        if isinstance(desde, (int, float)):
            factor_peso.append({'desdeKg': numero(desde), 'hastaKg': numero(hasta), 'sc': numero(sc), 'df': numero(df)})

    salida = {
        '_leeme': 'Modelo de tiempo estándar (MTM) de la zinguería. Generado con herramientas/importar-modelo-mtm.py '
                  'desde la planilla "Modelo MTM Zinguería" (hojas PARAMETROS_GLOBALES, ACTIVIDADES, MATERIALES, '
                  'TIPOS_PLIEGUE_Y_DETALLES y MTM_CODIGOS). Los valores se editan desde Producción > Parámetros de tiempos '
                  '(los cambios quedan guardados aparte y se pueden descartar). Estado: DATO = con fuente; [SUPUESTO] = '
                  'elegido por el analista; [A VALIDAR] = hay que cronometrarlo en planta.',
        'version': 'Modelo MTM Zinguería v1.0 (28/10/2026)',
        'tmuNota': '1 TMU = 0,00001 h = 0,036 s (MTM-1 Data Card, MTM ASSOCIATION e.V., 2019)',
        'calculados': [{'clave': c, 'formula': f} for c, f in {**CALCULADOS, **DESDE_CONFIG}.items()],
        'parametros': parametros + PARAMETROS_PROGRAMA,
        'productosSinPerfil': PRODUCTOS_SIN_PERFIL,
        'actividades': actividades,
        'materiales': materiales,
        'tiposPliegue': tipos_pliegue,
        'detalles': detalles,
        'factorPeso': factor_peso,
    }
    destino = Path(__file__).resolve().parent.parent / 'datos' / 'tiempos-mtm.json'
    destino.write_text(json.dumps(salida, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    print(f'Listo: {destino} ({len(parametros)} parámetros, {len(actividades)} actividades, {len(materiales)} materiales)')


if __name__ == '__main__':
    if len(sys.argv) != 2:
        sys.exit('Uso: python3 -I herramientas/importar-modelo-mtm.py "Modelo MTM Zinguería.xlsx"')
    main(sys.argv[1])
