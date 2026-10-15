"""Backup, restauración, rotación, exportación total y trabajos programados."""

import io
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import func, select

from circularis_rel.cli import migrar
from circularis_rel.config import Config, cargar_config
from circularis_rel.datos.db import BaseDatos
from circularis_rel.datos.modelos import Activo, Alerta, Base
from circularis_rel.datos.semillas import sembrar
from circularis_rel.servicios.demo import cargar_demo
from circularis_rel.servicios.exportacion import exportar_todo
from circularis_rel.tareas import backup, trabajos


def _instalacion(tmp_path, monkeypatch) -> tuple[BaseDatos, Config]:
    monkeypatch.setenv("CIRCULARIS_DATA_DIR", str(tmp_path / "datos"))
    monkeypatch.delenv("CIRCULARIS_DB_URL", raising=False)
    cfg = cargar_config()
    cfg.dir_datos.mkdir(parents=True)
    migrar(cfg.url_db)
    db = BaseDatos(cfg.url_db)
    with db.escritura() as s:
        sembrar(s)
    assert cargar_demo(db).aplicado
    return db, cfg


def test_exportacion_total(db):
    assert cargar_demo(db).aplicado
    with db.lectura() as s:
        wb = load_workbook(io.BytesIO(exportar_todo(s)), read_only=True)
    assert len(wb.sheetnames) == len(Base.metadata.tables)
    filas = list(wb["activo"].iter_rows(values_only=True))
    assert filas[0][0] == "id" and len(filas) == 4
    assert "hash_password" not in next(wb["usuario"].iter_rows(values_only=True))


def test_backup_restauracion_y_prueba(tmp_path, monkeypatch):
    db, cfg = _instalacion(tmp_path, monkeypatch)
    (cfg.dir_fotos / "2026" / "10").mkdir(parents=True)
    (cfg.dir_fotos / "2026" / "10" / "f.jpg").write_bytes(b"x")
    externa = tmp_path / "drive"
    archivo = backup.hacer_backup(cfg.url_db, cfg.dir_backups, cfg.dir_fotos, externa)
    assert archivo.exists() and (externa / archivo.name).exists()
    assert (cfg.dir_backups / archivo.name.replace("circularis_", "fotos_").replace(".db", ".zip")).exists()
    r = backup.probar_restauracion(cfg.dir_backups)
    assert r.ok and r.integridad == "ok"
    with db.escritura() as s:  # se rompe algo y se restaura
        s.delete(s.scalar(select(Activo).where(Activo.codigo == "DEMO-LAV-001")).posiciones[0])
    destino = tmp_path / "restaurada.db"
    backup.restaurar(archivo, f"sqlite:///{destino}", tmp_path / "fotos_rest")
    with BaseDatos(f"sqlite:///{destino}").lectura() as s:
        assert s.scalar(select(func.count()).select_from(Activo)) == 3
    assert (tmp_path / "fotos_rest" / "2026" / "10" / "f.jpg").exists()


def test_rotacion(tmp_path):
    d = tmp_path / "b"
    d.mkdir()
    for mes in (7, 8, 9, 10):
        for dia in (1, 10, 20):
            (d / f"circularis_2026{mes:02d}{dia:02d}_0300.db").write_bytes(b"")
    borrados = backup.rotar(d, diarios=2, mensuales=2)
    quedan = sorted(p.name for p in d.glob("*.db"))
    assert quedan == ["circularis_20260901_0300.db", "circularis_20261001_0300.db",
                      "circularis_20261010_0300.db", "circularis_20261020_0300.db"]
    assert len(borrados) == 8


def test_trabajos(tmp_path, monkeypatch):
    db, cfg = _instalacion(tmp_path, monkeypatch)
    r = trabajos.nocturno(db, cfg)
    assert "partes_estimados" in r
    assert trabajos.partes_faltantes(db, cfg) >= 0
    archivo = trabajos.hacer_backup(db, cfg)
    assert archivo.exists()
    assert trabajos.prueba_restauracion(db, cfg).ok
    with db.lectura() as s:
        assert s.scalar(select(func.count()).select_from(Alerta).where(Alerta.tipo == "BACKUP")) == 0
    assert isinstance(datetime.now(), datetime) and Path(cfg.dir_backups / "ultima_prueba_restauracion.json").exists()
