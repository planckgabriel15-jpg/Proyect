"""Fase 3: plan preventivo, calendario, OT, pruebas funcionales, sobre/sub-mantenimiento, repuestos, predictivo,
correo, reportes y formatos de la cátedra."""

import io
from dataclasses import replace
from datetime import date, datetime, time, timedelta
from zoneinfo import ZoneInfo

import pypdf
import pytest
from openpyxl import load_workbook
from sqlalchemy import func, select

from circularis_rel.asistente import rca
from circularis_rel.config import cargar_config
from circularis_rel.datos.enums import (
    EstadoEvento,
    EstadoOT,
    EstadoParte,
    OrigenEvento,
    ReglaPlan,
    Rol,
    TipoAnalisis,
    TipoFin,
    TipoOT,
)
from circularis_rel.datos.modelos import (
    ActivoTurno,
    Alerta,
    EventoFalla,
    IntervencionProgramada,
    OrdenTrabajo,
    ParteTurno,
    PosicionComponente,
    PrecioRepuesto,
    Repuesto,
    RepuestoActivo,
    TipoComponente,
    Turno,
    Usuario,
    VidaComponente,
)
from circularis_rel.servicios import (
    alertas,
    correo,
    formatos_catedra,
    ot,
    plan_preventivo,
    predictivo,
    recalculo,
    reporte,
    repuestos,
)
from circularis_rel.servicios import parametros as P
from circularis_rel.servicios.permisos import PermisoDenegado
from circularis_rel.servicios.registro import ErrorRegistro

TZ = ZoneInfo("America/Argentina/Cordoba")


def _ahora() -> datetime:
    return datetime.now(TZ)


@pytest.fixture
def continuo(db, planta):
    """Turno de 24 h todos los días en la máquina 1 con partes de los últimos 420 días (ritmo = 24 h/día)."""
    with db.escritura() as s:
        t = Turno(codigo="DEMO-24", nombre="Continuo", hora_inicio=time(0), hora_fin=time(0), duracion_h=24)
        s.add(t)
        s.flush()
        s.add(ActivoTurno(id_activo=planta.activo, id_turno=t.id, dias="1,2,3,4,5,6,7"))
        hoy = _ahora().date()
        for k in range(420):
            s.add(ParteTurno(id_activo=planta.activo, fecha=hoy - timedelta(days=k), id_turno=t.id, horas_turno=24,
                             estado=EstadoParte.TODO, horas_no_operadas=0, horas_operadas=24))
    return db


def _vida_en_servicio(s, id_posicion: int, horas: float) -> VidaComponente:
    v = VidaComponente(id_posicion=id_posicion, fecha_instalacion=_ahora() - timedelta(hours=horas))
    s.add(v)
    s.flush()
    return v


# ---------------------------------------------------------------------------
# Plan, calendario y OT
# ---------------------------------------------------------------------------


def test_aprobar_y_programar(continuo, planta):
    hoy = _ahora().date()
    with continuo.escritura() as s:
        _vida_en_servicio(s, planta.posicion, 240)  # 10 días de edad a 24 h/día
        plan = plan_preventivo.aprobar(s, planta.jefe, planta.posicion, ReglaPlan.FABRICANTE, 720)
        fechas, aviso = plan_preventivo.fechas_previstas(s, plan, hoy, TZ)
        assert aviso is None
        # El ritmo sale de los últimos 30 días de partes: 24 h por día. Faltan 480 h → unos 20 días; luego cada 30.
        assert (fechas[0] - hoy).days in (19, 20) and (fechas[1] - fechas[0]).days == 30
        r = plan_preventivo.programar(s, hoy, TZ)
        assert r == {"planes": 1, "intervenciones": len(fechas)}
        # Reprogramar no duplica.
        plan_preventivo.programar(s, hoy, TZ)
        assert s.scalar(select(func.count()).select_from(IntervencionProgramada)) == len(fechas)
        # Un plan nuevo deja al anterior fuera de vigencia y sus intervenciones sin OT se borran.
        plan_preventivo.aprobar(s, planta.jefe, planta.posicion, ReglaPlan.FABRICANTE, 360)
        plan_preventivo.programar(s, hoy, TZ)
        assert {iv.id_plan for iv in s.scalars(select(IntervencionProgramada))} == {plan_preventivo.plan_vigente(s, planta.posicion).id}


def test_aprobar_solo_el_jefe(db, planta):
    with db.escritura() as s:
        enc = Usuario(usuario="enc", nombre="Enc", rol=Rol.ENCARGADO, hash_password="x")
        s.add(enc)
        s.flush()
        with pytest.raises(PermisoDenegado):
            plan_preventivo.aprobar(s, enc.id, planta.posicion, ReglaPlan.FABRICANTE, 500)
        with pytest.raises(ErrorRegistro):
            plan_preventivo.aprobar(s, planta.jefe, planta.posicion, ReglaPlan.FABRICANTE, 0)


def test_ot_preventiva_con_anticipacion_y_cierre_por_cambio(continuo, planta):
    hoy = _ahora().date()
    with continuo.escritura() as s:
        _vida_en_servicio(s, planta.posicion, 700)  # ya pasó el intervalo: vence hoy
        plan_preventivo.aprobar(s, planta.jefe, planta.posicion, ReglaPlan.FABRICANTE, 600)
        plan_preventivo.programar(s, hoy, TZ)
        nuevas = plan_preventivo.generar_ot(s, hoy)
        assert len(nuevas) == 1 and nuevas[0].tipo == TipoOT.PREVENTIVA and nuevas[0].estado == EstadoOT.PLANIFICADA
        assert plan_preventivo.generar_ot(s, hoy) == []  # una OT abierta por plan
        cambio = ot.registrar_cambio(s, planta.jefe, planta.posicion, _ahora() - timedelta(minutes=1), hh=2)
        assert cambio.id == nuevas[0].id and cambio.estado == EstadoOT.CERRADA  # el cambio cierra la planificada
        assert s.scalar(select(func.count()).select_from(OrdenTrabajo)) == 1


def test_calendario_y_formatos(continuo, planta):
    hoy = _ahora().date()
    with continuo.escritura() as s:
        _vida_en_servicio(s, planta.posicion, 100)
        plan_preventivo.aprobar(s, planta.jefe, planta.posicion, ReglaPlan.FABRICANTE, 24 * 21)  # cada 21 días
        plan_preventivo.programar(s, hoy, TZ)
        filas = plan_preventivo.calendario(s, hoy)
        assert len(filas) == 1 and "Cambio de Cuchillas" in filas[0].tarea
        semanas = sorted(filas[0].semanas)
        assert all(b - a == 3 for a, b in zip(semanas, semanas[1:], strict=False))  # cada 3 semanas
        texto = "".join(p.extract_text() for p in pypdf.PdfReader(io.BytesIO(formatos_catedra.pdf_calendario(filas, hoy))).pages)
        assert "MANT/1CMP001" in texto and "MANT/2CMP001" in texto and "DEMO-AGR-001" in texto
        wb = load_workbook(io.BytesIO(formatos_catedra.excel_calendario(filas, hoy)))
        assert wb.sheetnames == ["MANT-1CMP001", "MANT-2CMP001"]
        nuevas = plan_preventivo.generar_ot(s, hoy + timedelta(days=30))
        texto = pypdf.PdfReader(io.BytesIO(formatos_catedra.pdf_ot_preventiva(s, nuevas[0].id, TZ))).pages[0].extract_text()
        assert "MANT/MPV001" in texto and "HERRAMIENTAS A UTILIZAR" in texto and "Reemplazar Cuchillas" in texto


# ---------------------------------------------------------------------------
# Fallas ocultas
# ---------------------------------------------------------------------------


def test_prueba_funcional(continuo, planta):
    hoy = _ahora().date()
    with continuo.escritura() as s:
        tipo = s.get(TipoComponente, planta.tipo2)
        tipo.falla_oculta = True
        pos = s.get(PosicionComponente, planta.posicion2)
        assert plan_preventivo.estado_prueba(s, pos, hoy, TZ)[0] == "GRIS"  # sin plan: "sin eventos" no es "todo bien"
        plan_preventivo.aprobar(s, planta.jefe, pos.id, ReglaPlan.PRUEBA_FUNCIONAL, 30 * 24)
        assert plan_preventivo.estado_prueba(s, pos, hoy, TZ)[0] == "ROJO"  # nunca se probó
        plan_preventivo.programar(s, hoy, TZ)
        prueba = plan_preventivo.generar_ot(s, hoy)[0]
        assert prueba.tipo == TipoOT.PRUEBA_FUNCIONAL
        plan_preventivo.cerrar_prueba_funcional(s, planta.jefe, prueba.id, _ahora() - timedelta(minutes=5), actuo=False)
        aviso = s.scalar(select(EventoFalla).where(EventoFalla.origen == OrigenEvento.PRUEBA_FUNCIONAL))
        assert aviso.estado == EstadoEvento.AVISADO and aviso.id_posicion == pos.id
        assert plan_preventivo.estado_prueba(s, pos, hoy, TZ)[0] == "VERDE"
        assert plan_preventivo.estado_prueba(s, pos, hoy + timedelta(days=26), TZ)[0] == "AMARILLO"
        assert plan_preventivo.estado_prueba(s, pos, hoy + timedelta(days=31), TZ)[0] == "ROJO"
        with pytest.raises(ErrorRegistro):
            plan_preventivo.cerrar_prueba_funcional(s, planta.jefe, prueba.id, _ahora(), actuo=True)  # ya cerrada


# ---------------------------------------------------------------------------
# Sobre y sub-mantenimiento
# ---------------------------------------------------------------------------


def _vidas_cerradas(s, id_posicion, id_modo, vidas):
    momento = _ahora() - timedelta(days=400)
    for edad, fin in vidas:
        s.add(VidaComponente(id_posicion=id_posicion, fecha_instalacion=momento, fecha_fin=momento + timedelta(hours=edad),
                             tipo_fin=fin, id_modo_falla=id_modo if fin == TipoFin.FALLA else None, validado=True, edad_horas=edad))
        momento += timedelta(hours=edad)
    s.flush()


def test_sub_mantenimiento(db, planta):
    with db.escritura() as s:
        _vidas_cerradas(s, planta.posicion, planta.modo, [(300, TipoFin.FALLA), (350, TipoFin.FALLA),
                                                          (500, TipoFin.SUSPENSION_PREVENTIVA), (420, TipoFin.FALLA)])
        plan = plan_preventivo.aprobar(s, planta.jefe, planta.posicion, ReglaPlan.FABRICANTE, 500)
        textos = plan_preventivo.diagnostico(s, s.get(PosicionComponente, planta.posicion), plan, [], TZ)
    assert len(textos) == 1 and textos[0].startswith("Sub-mantenimiento") and "75%" in textos[0]


def test_sobre_mantenimiento_con_beta_menor_a_uno(db, planta):
    from circularis_rel.datos.modelos import ResultadoWeibull

    with db.escritura() as s:
        plan = plan_preventivo.aprobar(s, planta.jefe, planta.posicion, ReglaPlan.FABRICANTE, 500)
        r = ResultadoWeibull(beta=0.8, eta=900)
        textos = plan_preventivo.diagnostico(s, s.get(PosicionComponente, planta.posicion), plan, [r], TZ)
    assert textos and textos[0].startswith("Sobre-mantenimiento")


# ---------------------------------------------------------------------------
# Repuestos
# ---------------------------------------------------------------------------


@pytest.fixture
def con_repuesto(continuo, planta):
    with continuo.escritura() as s:
        rep = Repuesto(codigo="DEMO-REP-1", descripcion="Juego de cuchillas", id_tipo_componente=planta.tipo,
                       stock_actual=0, lead_time_dias=10)
        s.add(rep)
        s.flush()
        s.add(PrecioRepuesto(id_repuesto=rep.id, fecha=date(2026, 1, 1), precio=100_000, moneda="ARS"))
        s.add(RepuestoActivo(id_repuesto=rep.id, id_activo=planta.activo, cantidad_instalada=1))
        _vidas_cerradas(s, planta.posicion, planta.modo, [(410, TipoFin.FALLA), (560, TipoFin.FALLA), (690, TipoFin.FALLA),
                                                          (780, TipoFin.FALLA), (920, TipoFin.FALLA)])
        _vida_en_servicio(s, planta.posicion, 300)
        recalculo.recalcular_todo(s, TZ)
        return rep.id


def test_demanda_y_stock_minimo(continuo, con_repuesto, planta):
    from circularis_rel.motor import kpis

    hoy = _ahora().date()
    with continuo.escritura() as s:
        e = next(x for x in repuestos.estado(s, hoy, TZ) if x.id == con_repuesto)
        ajustes = recalculo.ajustes_para_semaforo(s, planta.tipo, planta.activo, planta.posicion)
        vida = s.scalar(select(VidaComponente).where(VidaComponente.tipo_fin == TipoFin.EN_SERVICIO))
        esperado = kpis.prob_condicional_combinada(vida.edad_horas, 24 * 10, [(r.beta, r.eta, 0.0) for r in ajustes])
        assert e.demanda_lead == pytest.approx(esperado, rel=1e-6)
        assert e.stock_minimo >= 1 and e.bajo_minimo
        assert e.lote_economico is None and "costo de cada pedido (K)" in e.faltan
        P.fijar(s, "repuestos.costo_pedido_ars", 200, date(2000, 1, 2))
        P.fijar(s, "repuestos.tasa_almacenamiento_anual", 0.25, date(2000, 1, 2))
        e = next(x for x in repuestos.estado(s, hoy, TZ) if x.id == con_repuesto)
        assert e.lote_economico == pytest.approx((2 * 200 * e.demanda_anual / (0.25 * 100_000)) ** 0.5)
        assert e.stock_proteccion > 0
        assert repuestos.alertas_stock(s, hoy, TZ) == 1
        assert s.scalar(select(Alerta).where(Alerta.tipo == "STOCK", Alerta.resuelta_en.is_(None))) is not None
        repuestos.ajustar_stock(s, con_repuesto, 10, "conteo")
        repuestos.alertas_stock(s, hoy, TZ)
        assert s.scalar(select(Alerta).where(Alerta.tipo == "STOCK", Alerta.resuelta_en.is_(None))) is None
        with pytest.raises(ValueError):
            repuestos.ajustar_stock(s, con_repuesto, 5, " ")


def test_repuesto_sin_tipo(db, planta):
    with db.escritura() as s:
        s.add(Repuesto(codigo="R-SUELTO", descripcion="Suelto"))
        s.flush()
        e = repuestos.estado(s, _ahora().date(), TZ)[0]
    assert e.demanda_lead is None and e.stock_minimo is None and "lead time" in e.faltan


# ---------------------------------------------------------------------------
# Predictivo
# ---------------------------------------------------------------------------


def test_predictivo_ultrasonido(db, planta):
    with db.escritura() as s:
        p = predictivo.crear_parametro(s, planta.jefe, planta.activo, "Ultrasonido", "Rulemán", "dB", linea_base=30,
                                       plantilla=predictivo.PLANTILLA_ULTRASONIDO)
        assert (p.limite_alerta, p.limite_accion) == (38, 46)  # +8 dB y +16 dB de la tabla de la UNIDAD 3
        for nombre in ("B", "C"):
            predictivo.crear_parametro(s, planta.jefe, planta.activo, "Ultrasonido", nombre, "dB")
        with pytest.raises(ErrorRegistro, match="hasta 3"):
            predictivo.crear_parametro(s, planta.jefe, planta.activo, "Ultrasonido", "D", "dB")
        assert predictivo.registrar(s, planta.jefe, p.id, _ahora() - timedelta(days=3), 33).nivel == "NORMAL"
        ev = predictivo.registrar(s, planta.jefe, p.id, _ahora() - timedelta(days=2), 41)
        assert ev.nivel == "ALERTA" and "Vigilar" in ev.texto and ev.ot is None
        ev = predictivo.registrar(s, planta.jefe, p.id, _ahora() - timedelta(days=1), 47)
        assert ev.nivel == "ACCION" and ev.ot is not None and ev.ot.tipo == TipoOT.INSPECCION
        assert predictivo.registrar(s, planta.jefe, p.id, _ahora(), 48).ot.id == ev.ot.id  # una OT mientras siga abierta
        t = predictivo.tendencia_parametro(s, p)
        assert t is not None and t.pendiente_por_dia > 0 and t.dias_hasta_limite == 0.0


def test_predictivo_parametro_que_empeora_al_bajar(db, planta):
    with db.escritura() as s:
        p = predictivo.crear_parametro(s, planta.jefe, planta.activo, "Eléctrica", "Aislación", "MΩ", linea_base=500,
                                       limite_alerta=100, limite_accion=50)
        assert predictivo.registrar(s, planta.jefe, p.id, _ahora(), 80).nivel == "ALERTA"
        with pytest.raises(ErrorRegistro, match="antes"):
            predictivo.crear_parametro(s, planta.jefe, planta.activo, "Eléctrica", "Otra", "MΩ", linea_base=500,
                                       limite_alerta=40, limite_accion=50)


# ---------------------------------------------------------------------------
# Alertas diarias, correo y reportes
# ---------------------------------------------------------------------------


def test_eventos_sin_validar(db, planta):
    with db.escritura() as s:
        falla = _ahora() - timedelta(days=5)
        ev = EventoFalla(id_activo=planta.activo, id_posicion=planta.posicion, fecha_hora_falla=falla, id_modo=planta.modo,
                         id_causa=planta.causa, inicio_reparacion=falla, fin_reparacion=falla + timedelta(hours=1),
                         estado=EstadoEvento.CERRADO)
        s.add(ev)
        s.flush()
        assert alertas.eventos_sin_validar(s) == 0  # recién cerrado
        ev.actualizado_en = _ahora() - timedelta(hours=60)
        s.flush()
        from sqlalchemy import update

        s.execute(update(EventoFalla).where(EventoFalla.id == ev.id).values(actualizado_en=_ahora() - timedelta(hours=60)))
        assert alertas.eventos_sin_validar(s) == 1


class _SMTPFalso:
    enviados: list = []

    def __init__(self, *a, **k):
        pass

    def __enter__(self):
        return self

    def __exit__(self, *a):
        return False

    def starttls(self, **k):
        pass

    def login(self, *a):
        pass

    def send_message(self, m):
        _SMTPFalso.enviados.append(m)


def test_correo(db, planta, monkeypatch):
    cfg = cargar_config()
    with db.escritura() as s:
        s.get(Usuario, planta.jefe).email = "jefe@planta.local"
        alertas.emitir(s, "RIESGO", "alta", "AGR · CUCH en rojo", "riesgo:1", destinatarios="JEFE")
        alertas.emitir(s, "RIESGO_ORIENTATIVO", "baja", "orientativa", "riesgo_orientativo:1")
        alertas.emitir(s, "STOCK", "media", "stock bajo", "stock:1", destinatarios="JEFE")
        assert correo.enviar_alertas(s, cfg, inmediatas=True) == 0  # sin SMTP no se manda nada
        monkeypatch.setattr(correo.smtplib, "SMTP", _SMTPFalso)
        _SMTPFalso.enviados = []
        cfg_smtp = replace(cfg, smtp_servidor="smtp.local", smtp_remitente="circularis@planta.local")
        assert correo.enviar_alertas(s, cfg_smtp, inmediatas=True) == 1
        assert correo.enviar_alertas(s, cfg_smtp, inmediatas=True) == 0  # ya marcada como enviada
        assert correo.enviar_alertas(s, cfg_smtp, inmediatas=False) == 1  # el stock va en el resumen diario
        assert [m["To"] for m in _SMTPFalso.enviados] == ["jefe@planta.local", "jefe@planta.local"]
        orientativa = s.scalar(select(Alerta).where(Alerta.tipo == "RIESGO_ORIENTATIVO"))
        assert orientativa.enviada_mail_en is None  # nunca se manda


def test_reporte_pdf(db, planta):
    from circularis_rel.servicios.demo import cargar_demo

    assert cargar_demo(db).aplicado
    hoy = _ahora().date()
    with db.lectura() as s:
        datos = reporte.pdf(s, *reporte.mes_anterior(hoy), TZ)
    texto = "".join(p.extract_text() for p in pypdf.PdfReader(io.BytesIO(datos)).pages)
    for parte in ("DATOS FICTICIOS", "Resumen de la línea", "Estado de las máquinas", "Tablero de comando", "Calidad de los datos"):
        assert parte in texto
    assert reporte.mes_anterior(date(2026, 1, 15)) == (date(2025, 12, 1), date(2025, 12, 31))


def test_pdf_predictivo(db, planta):
    with db.escritura() as s:
        p = predictivo.crear_parametro(s, planta.jefe, planta.activo, "Vibraciones", "Vibración LA", "mm/s", linea_base=2,
                                       limite_alerta=4.5, limite_accion=7.1)
        predictivo.registrar(s, planta.jefe, p.id, _ahora(), 3.2, "Vibrómetro")
        texto = pypdf.PdfReader(io.BytesIO(formatos_catedra.pdf_predictivo(s, planta.activo, "Vibraciones", TZ))).pages[0].extract_text()
    assert "MANT/PMRED001" in texto and "Vibración LA" in texto and "Vibrómetro" in texto


# ---------------------------------------------------------------------------
# Ishikawa, 5W+1H y hoja de verificación
# ---------------------------------------------------------------------------


def test_otras_herramientas_de_analisis(db, planta):
    with db.escritura() as s:
        a = rca.crear(s, planta.jefe, "Rulemanes", id_activo=planta.activo, tipo=TipoAnalisis.ISHIKAWA_5M)
        with pytest.raises(ErrorRegistro):
            rca.guardar_contenido(s, planta.jefe, a.id, {"causas": {}})
        rca.guardar_contenido(s, planta.jefe, a.id, {"causas": {"Métodos": ["Sin plan de lubricación", " "]},
                                                     "principal": "Sin plan de lubricación"})
        assert a.contenido["causas"]["Métodos"] == ["Sin plan de lubricación"] and a.causa_raiz == "Sin plan de lubricación"
        svg = rca.ishikawa_svg("Rulemanes", a.contenido["causas"])
        assert svg.startswith("<svg") and "Medio ambiente" in svg and "Sin plan de lubricación" in svg
        w = rca.crear(s, planta.jefe, "Corte", id_activo=planta.activo, tipo=TipoAnalisis.CINCO_W_1H)
        rca.guardar_contenido(s, planta.jefe, w.id, {"que": "Se corta la correa", "por_que": "Polea desalineada"})
        assert w.causa_raiz == "Polea desalineada"
        h = rca.crear(s, planta.jefe, "Conteo", id_activo=planta.activo, tipo=TipoAnalisis.HOJA_VERIFICACION)
        rca.guardar_contenido(s, planta.jefe, h.id, {"items": [{"categoria": "Ruido", "conteo": 2}, {"categoria": "Fuga", "conteo": 5}]})
        assert [i["categoria"] for i in h.contenido["items"]] == ["Fuga", "Ruido"]
        rca.guardar_accion(s, planta.jefe, h.id, causa_raiz="Juntas", accion="Cambiar juntas", responsable="Enc",
                           fecha_compromiso=None)
        assert h.accion_correctiva == "Cambiar juntas"
