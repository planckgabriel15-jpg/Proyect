"""Importador: validación por fila, simular, todo o nada, idempotencia y triage del histórico."""

from datetime import date
from pathlib import Path

import pytest
from openpyxl import load_workbook
from sqlalchemy import func, select

from circularis_rel.datos.enums import ClaseCriticidad, ClaseTriage
from circularis_rel.datos.modelos import (
    Activo,
    AmfeItem,
    Auditoria,
    Catalogo,
    CatalogoAplica,
    Importacion,
    PosicionComponente,
    PrecioRepuesto,
    VidaHistorica,
)
from circularis_rel.servicios.importador.aplicar import importar
from circularis_rel.servicios.importador.esquema import HOJA_HISTORICO, HOJAS_CARGA_INICIAL
from circularis_rel.servicios.importador.plantillas import escribir_libro, plantilla_carga_inicial
from circularis_rel.servicios.importador.triage import parece_reconstruido

BASE = {
    "lineas": [{"codigo": "DEMO-L", "nombre": "Línea DEMO"}],
    "turnos": [{"codigo": "M", "nombre": "Mañana", "hora_inicio": "06:00", "hora_fin": "14:00", "duracion_h": 8}],
    "tipos_componente": [{"codigo": "DEMO-CUCH", "nombre": "Cuchillas"}, {"codigo": "DEMO-ROD", "nombre": "Rulemanes"}],
    "maquinas": [{"codigo": "DEMO-AGR-001", "nombre": "Agrumadora DEMO", "linea": "DEMO-L",
                  "etapa_proceso": "PREPROCESADO", "horas_promedio_dia": "8", "fecha_instalacion": "15/03/2022"}],
    "maquina_turnos": [{"maquina": "DEMO-AGR-001", "turno": "M", "dias": "1,2,3,4,5"}],
    "posiciones": [{"maquina": "DEMO-AGR-001", "codigo": "CUCH", "tipo_componente": "DEMO-CUCH", "critica": "SI"},
                   {"maquina": "DEMO-AGR-001", "codigo": "ROD", "tipo_componente": "DEMO-ROD"}],
    "modos": [{"codigo": "DEMO-FILO", "descripcion": "Pérdida de filo", "tipos_componente": "DEMO-CUCH"}],
    "repuestos": [{"codigo": "DEMO-R1", "descripcion": "Juego de cuchillas", "tipo_componente": "DEMO-CUCH",
                   "precio": "1.495,50", "fecha_precio": "01/10/2026"}],
    "repuestos_por_maquina": [{"maquina": "DEMO-AGR-001", "repuesto": "DEMO-R1", "cantidad_instalada": 1}],
    "criticidad": [{"maquina": "DEMO-AGR-001", "uso": 5, "alternativa": 5, "influencia": 5, "calidad": 5,
                    "ambiente": 3, "seguridad": 5}],
    "amfe": [{"maquina": "DEMO-AGR-001", "tipo_componente": "DEMO-CUCH", "modo": "DEMO-FILO", "g": 6, "o": 5, "d": 4}],
}


def libro(tmp_path: Path, datos: dict, nombre="carga.xlsx") -> Path:
    return escribir_libro(tmp_path / nombre, HOJAS_CARGA_INICIAL, datos)


def contar(db, modelo):
    with db.lectura() as s:
        return s.scalar(select(func.count()).select_from(modelo))


def test_plantilla_vacia_tiene_todas_las_hojas_y_listas(tmp_path):
    ruta = plantilla_carga_inicial(tmp_path / "vacia.xlsx")
    wb = load_workbook(ruta)
    assert [h.nombre for h in HOJAS_CARGA_INICIAL] == wb.sheetnames[1:]
    ws = wb["maquinas"]
    assert ws.cell(1, 1).value == "codigo"
    assert "obligatoria" in ws.cell(2, 1).value
    assert any("PREPROCESADO" in dv.formula1 for dv in ws.data_validations.dataValidation)


def test_simular_no_graba(db, tmp_path):
    r = importar(db, libro(tmp_path, BASE), simular=True)
    assert r.ok and not r.aplicado
    assert r.hojas["maquinas"].creados == 1
    assert contar(db, Activo) == 0


def test_aplicar_y_reimportar_no_duplica(db, tmp_path):
    ruta = libro(tmp_path, BASE)
    r1 = importar(db, ruta, simular=False, usuario_id=None)
    assert r1.aplicado, r1.texto()
    r2 = importar(db, ruta, simular=False)
    assert r2.aplicado
    assert all(h.creados == 0 and h.actualizados == 0 for h in r2.hojas.values()), r2.texto()
    assert contar(db, Activo) == 1 and contar(db, PosicionComponente) == 2 and contar(db, AmfeItem) == 1
    assert contar(db, Importacion) == 2
    with db.lectura() as s:
        a = s.scalar(select(Activo))
        assert a.fecha_instalacion == date(2022, 3, 15)
        assert a.criticidad_puntaje == 28 and a.criticidad_nivel == "Alta" and a.criticidad_clase == ClaseCriticidad.A
        assert s.scalar(select(PrecioRepuesto.precio)) == pytest.approx(1495.50)
        acciones = set(s.scalars(select(Auditoria.accion).where(Auditoria.tabla == "activo")))
        assert acciones == {"IMPORTACION"}


def test_actualizacion_cambia_solo_lo_distinto(db, tmp_path):
    importar(db, libro(tmp_path, BASE), simular=False)
    datos = {**BASE, "maquinas": [dict(BASE["maquinas"][0], horas_promedio_dia=6)]}
    r = importar(db, libro(tmp_path, datos, "v2.xlsx"), simular=False)
    assert r.hojas["maquinas"].actualizados == 1
    with db.lectura() as s:
        assert s.scalar(select(Activo.horas_promedio_dia)) == 6


def test_errores_por_fila_y_todo_o_nada(db, tmp_path):
    datos = {**BASE,
             "maquinas": BASE["maquinas"] + [
                 {"codigo": "DEMO-X", "nombre": "Sin línea", "linea": "NO-EXISTE", "etapa_proceso": "LAVADO"},
                 {"codigo": "DEMO-Y", "nombre": "Etapa mala", "linea": "DEMO-L", "etapa_proceso": "INVENTADA"},
                 {"codigo": "DEMO-Z", "nombre": "Horas malas", "linea": "DEMO-L", "etapa_proceso": "LAVADO",
                  "horas_promedio_dia": 30}],
             "posiciones": BASE["posiciones"] + [{"maquina": "DEMO-AGR-001", "codigo": "P3"}]}
    r = importar(db, libro(tmp_path, datos), simular=False)
    assert not r.aplicado
    textos = [str(e) for e in r.errores]
    assert any("fila 4" in t and "no existe una línea" in t for t in textos)
    assert any("fila 5" in t and "etapa_proceso" in t for t in textos)
    assert any("fila 6" in t and "máximo" in t for t in textos)
    assert any("posiciones" in t and "tipo_componente" in t and "obligatoria" in t for t in textos)
    assert contar(db, Activo) == 0  # nada grabado


def test_modo_con_tipo_inexistente(db, tmp_path):
    datos = {**BASE, "modos": [{"codigo": "M", "descripcion": "x", "tipos_componente": "NO-EXISTE"}]}
    r = importar(db, libro(tmp_path, datos), simular=True)
    assert any("NO-EXISTE" in str(e) for e in r.errores)


def test_modos_reemplazan_tipos_aplicables(db, tmp_path):
    importar(db, libro(tmp_path, BASE), simular=False)
    datos = {**BASE, "modos": [{"codigo": "DEMO-FILO", "descripcion": "Pérdida de filo",
                                "tipos_componente": "DEMO-CUCH, DEMO-ROD"}],
             "amfe": []}
    importar(db, libro(tmp_path, datos, "v2.xlsx"), simular=False)
    with db.lectura() as s:
        modo = s.scalar(select(Catalogo).where(Catalogo.codigo == "DEMO-FILO"))
        assert s.scalar(select(func.count()).select_from(CatalogoAplica).where(CatalogoAplica.id_catalogo == modo.id)) == 2


def test_amfe_con_modo_no_asignado_al_tipo(db, tmp_path):
    datos = {**BASE, "amfe": [{"tipo_componente": "DEMO-ROD", "modo": "DEMO-FILO", "g": 1, "o": 1, "d": 1}]}
    r = importar(db, libro(tmp_path, datos), simular=True)
    assert any("no está asignado" in str(e) for e in r.errores)


def test_criticidad_con_valor_fuera_de_escala(db, tmp_path):
    datos = {**BASE, "criticidad": [dict(BASE["criticidad"][0], uso=4)]}
    r = importar(db, libro(tmp_path, datos), simular=True)
    assert any("escala" in str(e) for e in r.errores)


def test_beta_ref_sin_fuente(db, tmp_path):
    datos = {**BASE, "tipos_componente": [{"codigo": "DEMO-CUCH", "nombre": "Cuchillas", "beta_ref": 2}]}
    r = importar(db, libro(tmp_path, datos), simular=True)
    assert any("fuente" in str(e) for e in r.errores)


def test_precio_sin_fecha(db, tmp_path):
    datos = {**BASE, "repuestos": [{"codigo": "DEMO-R1", "descripcion": "x", "precio": 10}]}
    r = importar(db, libro(tmp_path, datos), simular=True)
    assert any("fecha" in str(e) for e in r.errores)


def test_columna_desconocida(db, tmp_path):
    ruta = libro(tmp_path, BASE)
    wb = load_workbook(ruta)
    wb["lineas"].cell(1, 3, "columna_rara")
    wb.save(ruta)
    r = importar(db, ruta, simular=True)
    assert any("columna_rara" in str(e) for e in r.errores)


# --- Histórico -------------------------------------------------------------

def _historico(tmp_path, filas):
    return escribir_libro(tmp_path / "hist.xlsx", (HOJA_HISTORICO,), {"historico": filas})


def test_triage_clases(db, tmp_path):
    importar(db, libro(tmp_path, BASE), simular=False)
    filas = [
        # A: fecha real, posición, modo, respaldo
        {"maquina": "DEMO-AGR-001", "posicion": "CUCH", "tipo_fin": "FALLA", "modo": "DEMO-FILO",
         "fecha_fin": "10/05/2025", "fuente": "OT", "respaldo": "OT 123"},
        # B: posición y rango de fechas
        {"maquina": "DEMO-AGR-001", "posicion": "CUCH", "tipo_fin": "FALLA", "modo": "DEMO-FILO",
         "fecha_fin_desde": "01/06/2025", "fecha_fin_hasta": "30/06/2025", "fuente": "Factura de cuchillas"},
        # C: sin posición
        {"maquina": "DEMO-AGR-001", "tipo_fin": "FALLA", "fecha_fin": "01/07/2025", "fuente": "Memoria del encargado"},
    ]
    r = importar(db, _historico(tmp_path, filas), simular=False, historico=True)
    assert r.aplicado, r.texto()
    with db.lectura() as s:
        clases = list(s.scalars(select(VidaHistorica.clase_triage).order_by(VidaHistorica.id)))
    assert clases == [ClaseTriage.A, ClaseTriage.B, ClaseTriage.C]
    r2 = importar(db, _historico(tmp_path, filas), simular=False, historico=True)
    assert r2.hojas["historico"].sin_cambios == 3


def test_triage_clase_a_sin_respaldo_es_error(db, tmp_path):
    importar(db, libro(tmp_path, BASE), simular=False)
    filas = [{"maquina": "DEMO-AGR-001", "posicion": "CUCH", "tipo_fin": "FALLA", "modo": "DEMO-FILO",
              "fecha_fin": "10/05/2025", "fuente": "memoria", "clase": "A"}]
    r = importar(db, _historico(tmp_path, filas), simular=True, historico=True)
    assert any("respaldo" in str(e) for e in r.errores)


def test_triage_rango_incompleto(db, tmp_path):
    importar(db, libro(tmp_path, BASE), simular=False)
    filas = [{"maquina": "DEMO-AGR-001", "posicion": "CUCH", "tipo_fin": "FALLA",
              "fecha_fin_desde": "01/06/2025", "fuente": "x"}]
    r = importar(db, _historico(tmp_path, filas), simular=True, historico=True)
    assert any("desde' y 'hasta'" in str(e) for e in r.errores)


def test_espaciado_regular_se_advierte(db, tmp_path):
    """Plan §11: 90, 120, 150, 180… días son típicos de valores reconstruidos (TRI-001 del TP)."""
    importar(db, libro(tmp_path, BASE), simular=False)
    filas = [{"maquina": "DEMO-AGR-001", "posicion": "CUCH", "tipo_fin": "SUSPENSION_OTRA",
              "fecha_fin": f"{d:02d}/0{m}/2025", "fuente": "planilla"} for d, m in ((1, 1), (31, 1), (2, 3), (1, 4))]
    r = importar(db, _historico(tmp_path, filas), simular=True, historico=True)
    assert any("reconstruidos" in str(a) for a in r.advertencias)


def test_parece_reconstruido():
    assert parece_reconstruido([date(2025, 1, 1), date(2025, 4, 1), date(2025, 7, 1), date(2025, 10, 1)]) is False  # 90,91,92
    assert parece_reconstruido([date(2025, 1, 1), date(2025, 3, 2), date(2025, 5, 1), date(2025, 7, 30)])  # 60,60,90
    assert not parece_reconstruido([date(2025, 1, 1), date(2025, 1, 20)])


def test_borrador_tp1_importa_sin_errores(db):
    ruta = Path(__file__).resolve().parents[2] / "docs" / "plantillas" / "carga_inicial_borrador_TP1.xlsx"
    r = importar(db, ruta, simular=False)
    assert r.aplicado, r.texto()
    with db.lectura() as s:
        clases = dict(s.execute(select(Activo.codigo, Activo.criticidad_clase)).all())
        origenes = set(s.scalars(select(AmfeItem.origen)))
    assert len(clases) == 10
    assert clases["EXT-001"] == "A" and clases["SEC-001"] == "A"  # 28 y 26 puntos: Alta
    assert clases["AGR-001"] == "B" and clases["LAV-001"] == "B"  # 23 y 24: Media
    assert clases["TRI-001"] == "C" and clases["ENF-001"] == "C"  # 17 Baja, 14 Muy baja
    assert origenes == {"TP1_ILUSTRATIVO"}
