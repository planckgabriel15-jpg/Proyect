"""Generación de plantillas Excel vacías (o con filas precargadas) a partir del esquema."""

from __future__ import annotations

from collections.abc import Iterable, Mapping
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from circularis_rel.servicios.importador.esquema import HOJA_HISTORICO, HOJAS_CARGA_INICIAL, Hoja

FILA_ENCABEZADO = 1
FILA_AYUDA = 2
PRIMERA_FILA_DATOS = 3
FILAS_CON_VALIDACION = 500

_RELLENO_ENCABEZADO = PatternFill("solid", fgColor="1F4E78")
_RELLENO_OBLIGATORIA = PatternFill("solid", fgColor="C00000")
_RELLENO_AYUDA = PatternFill("solid", fgColor="F2F2F2")


def _texto_ayuda(col) -> str:
    partes = []
    if col.obligatoria:
        partes.append("* obligatoria")
    if col.ayuda:
        partes.append(col.ayuda)
    if col.opciones and col.tipo in ("lista", "bool"):
        partes.append("opciones: " + ", ".join(col.opciones))
    return ". ".join(partes)


def escribir_hoja(wb: Workbook, hoja: Hoja, filas: Iterable[Mapping] = ()) -> None:
    ws = wb.create_sheet(hoja.nombre)
    for i, col in enumerate(hoja.columnas, start=1):
        celda = ws.cell(row=FILA_ENCABEZADO, column=i, value=col.nombre)
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = _RELLENO_OBLIGATORIA if col.obligatoria else _RELLENO_ENCABEZADO
        ayuda = ws.cell(row=FILA_AYUDA, column=i, value=_texto_ayuda(col))
        ayuda.font = Font(italic=True, color="595959", size=9)
        ayuda.fill = _RELLENO_AYUDA
        ayuda.alignment = Alignment(wrap_text=True, vertical="top")
        ws.column_dimensions[get_column_letter(i)].width = max(14, min(40, len(col.nombre) + 6))
        if col.opciones and col.tipo in ("lista", "bool"):
            dv = DataValidation(type="list", formula1='"' + ",".join(col.opciones) + '"', allow_blank=True,
                                showErrorMessage=True, errorTitle="Valor no válido",
                                error="Elegí un valor de la lista")
            letra = get_column_letter(i)
            dv.add(f"{letra}{PRIMERA_FILA_DATOS}:{letra}{PRIMERA_FILA_DATOS + FILAS_CON_VALIDACION}")
            ws.add_data_validation(dv)
    ws.row_dimensions[FILA_AYUDA].height = 45
    ws.freeze_panes = ws.cell(row=PRIMERA_FILA_DATOS, column=1)
    for r, fila in enumerate(filas, start=PRIMERA_FILA_DATOS):
        for i, col in enumerate(hoja.columnas, start=1):
            valor = fila.get(col.nombre)
            if isinstance(valor, bool):
                valor = "SI" if valor else "NO"
            if valor is not None:
                ws.cell(row=r, column=i, value=valor)


def escribir_libro(ruta: Path, hojas: Iterable[Hoja], datos: Mapping[str, Iterable[Mapping]] | None = None,
                   nota: str | None = None) -> Path:
    datos = datos or {}
    wb = Workbook()
    portada = wb.active
    portada.title = "LEEME"
    lineas = [
        "Circularis · Plantilla de importación",
        "Fila 1: encabezados (no cambiar). Fila 2: ayuda (se ignora). Los datos empiezan en la fila 3.",
        "Encabezados en rojo: columnas obligatorias. Las listas desplegables muestran los valores aceptados.",
        "Re-importar el mismo libro no duplica: la clave es el código.",
        "Primero usar 'Simular' en el sistema: muestra los errores fila por fila sin grabar nada.",
    ]
    if nota:
        lineas += ["", nota]
    for i, t in enumerate(lineas, start=1):
        portada.cell(row=i, column=1, value=t).font = Font(bold=(i == 1), size=12 if i == 1 else 10)
    portada.column_dimensions["A"].width = 120
    for hoja in hojas:
        escribir_hoja(wb, hoja, datos.get(hoja.nombre, ()))
    ruta.parent.mkdir(parents=True, exist_ok=True)
    wb.save(ruta)
    return ruta


def plantilla_carga_inicial(ruta: Path) -> Path:
    return escribir_libro(ruta, HOJAS_CARGA_INICIAL)


def plantilla_historico(ruta: Path) -> Path:
    return escribir_libro(ruta, (HOJA_HISTORICO,), nota=(
        "Triage: A = confiable (fecha real, componente y modo, con respaldo); "
        "B = parcial (componente claro, fecha aproximada: completar desde/hasta); "
        "C = no usable (solo Pareto cualitativo). Nunca se importa directo al cálculo."))
