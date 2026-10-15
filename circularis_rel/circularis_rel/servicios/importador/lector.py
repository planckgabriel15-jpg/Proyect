"""Lectura de un libro Excel según el esquema, con errores por fila y columna."""

from __future__ import annotations

from dataclasses import dataclass, field
from datetime import date, datetime, time
from pathlib import Path

from openpyxl import load_workbook

from circularis_rel.servicios.importador.esquema import Columna, Hoja
from circularis_rel.servicios.importador.plantillas import FILA_ENCABEZADO, PRIMERA_FILA_DATOS


@dataclass
class Mensaje:
    hoja: str
    fila: int | None
    columna: str | None
    texto: str

    def __str__(self) -> str:
        donde = self.hoja
        if self.fila:
            donde += f", fila {self.fila}"
        if self.columna:
            donde += f", columna '{self.columna}'"
        return f"{donde}: {self.texto}"


@dataclass
class FilaLeida:
    numero: int
    valores: dict


@dataclass
class HojaLeida:
    hoja: Hoja
    filas: list[FilaLeida] = field(default_factory=list)


class ValorInvalido(ValueError):
    pass


def _numero(valor) -> float:
    if isinstance(valor, bool):
        raise ValorInvalido("se esperaba un número")
    if isinstance(valor, int | float):
        return float(valor)
    texto = str(valor).strip().replace(" ", "")
    if "," in texto and "." in texto:  # 1.495,50 -> 1495.50
        texto = texto.replace(".", "").replace(",", ".")
    else:
        texto = texto.replace(",", ".")
    try:
        return float(texto)
    except ValueError as e:
        raise ValorInvalido(f"'{valor}' no es un número") from e


def convertir(valor, col: Columna):
    if valor is None or (isinstance(valor, str) and not valor.strip()):
        return None
    t = col.tipo
    if t == "texto":
        if isinstance(valor, float) and valor.is_integer():
            valor = int(valor)
        return str(valor).strip()
    if t in ("real", "entero"):
        n = _numero(valor)
        if t == "entero":
            if not float(n).is_integer():
                raise ValorInvalido(f"'{valor}' no es un número entero")
            n = int(n)
        if col.minimo is not None and n < col.minimo:
            raise ValorInvalido(f"{n} es menor que el mínimo {col.minimo}")
        if col.maximo is not None and n > col.maximo:
            raise ValorInvalido(f"{n} es mayor que el máximo {col.maximo}")
        return n
    if t == "fecha":
        if isinstance(valor, datetime):
            return valor.date()
        if isinstance(valor, date):
            return valor
        texto = str(valor).strip()
        for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y", "%d/%m/%y"):
            try:
                return datetime.strptime(texto, fmt).date()
            except ValueError:
                continue
        raise ValorInvalido(f"'{valor}' no es una fecha DD/MM/AAAA")
    if t == "hora":
        if isinstance(valor, time):
            return valor.replace(second=0, microsecond=0)
        if isinstance(valor, datetime):
            return valor.time().replace(second=0, microsecond=0)
        if isinstance(valor, float) and 0 <= valor < 1:  # fracción de día de Excel
            minutos = round(valor * 24 * 60)
            return time(minutos // 60, minutos % 60)
        texto = str(valor).strip()
        for fmt in ("%H:%M", "%H:%M:%S", "%H.%M"):
            try:
                return datetime.strptime(texto, fmt).time()
            except ValueError:
                continue
        raise ValorInvalido(f"'{valor}' no es una hora HH:MM")
    if t == "bool":
        texto = str(valor).strip().upper()
        if texto in ("SI", "SÍ", "S", "1", "TRUE", "VERDADERO", "X"):
            return True
        if texto in ("NO", "N", "0", "FALSE", "FALSO"):
            return False
        raise ValorInvalido(f"'{valor}' no es SI o NO")
    if t == "lista":
        texto = str(valor).strip()
        for op in col.opciones:
            if texto.upper() == op.upper():
                return op
        raise ValorInvalido(f"'{valor}' no está en la lista: {', '.join(col.opciones)}")
    if t == "codigos":
        codigos = [c.strip() for c in str(valor).replace(";", ",").split(",") if c.strip()]
        return codigos or None
    raise AssertionError(f"tipo de columna desconocido: {t}")


def leer_libro(ruta: Path, hojas: tuple[Hoja, ...], errores: list[Mensaje]) -> dict[str, HojaLeida]:
    """Lee las hojas presentes. Una hoja ausente no es error (carga parcial)."""
    wb = load_workbook(ruta, data_only=True, read_only=True)
    salida: dict[str, HojaLeida] = {}
    for hoja in hojas:
        if hoja.nombre not in wb.sheetnames:
            continue
        ws = wb[hoja.nombre]
        filas = list(ws.iter_rows(values_only=True))
        if not filas:
            continue
        encabezados = [str(h).strip() if h is not None else "" for h in filas[FILA_ENCABEZADO - 1]]
        faltan = [c.nombre for c in hoja.columnas if c.obligatoria and c.nombre not in encabezados]
        if faltan:
            errores.append(Mensaje(hoja.nombre, FILA_ENCABEZADO, None,
                                   f"faltan columnas obligatorias: {', '.join(faltan)}"))
            continue
        desconocidas = [h for h in encabezados if h and h not in {c.nombre for c in hoja.columnas}]
        if desconocidas:
            errores.append(Mensaje(hoja.nombre, FILA_ENCABEZADO, None,
                                   f"columnas desconocidas (revisar el nombre): {', '.join(desconocidas)}"))
            continue
        indice = {h: i for i, h in enumerate(encabezados) if h}
        leida = HojaLeida(hoja)
        for n, fila in enumerate(filas[PRIMERA_FILA_DATOS - 1:], start=PRIMERA_FILA_DATOS):
            if fila is None or all(v is None or (isinstance(v, str) and not v.strip()) for v in fila):
                continue
            valores: dict = {}
            ok = True
            for col in hoja.columnas:
                crudo = fila[indice[col.nombre]] if col.nombre in indice and indice[col.nombre] < len(fila) else None
                try:
                    valores[col.nombre] = convertir(crudo, col)
                except ValorInvalido as e:
                    errores.append(Mensaje(hoja.nombre, n, col.nombre, str(e)))
                    ok = False
                    continue
                if col.obligatoria and valores[col.nombre] is None:
                    errores.append(Mensaje(hoja.nombre, n, col.nombre, "es obligatoria"))
                    ok = False
            if ok:
                leida.filas.append(FilaLeida(n, valores))
        salida[hoja.nombre] = leida
    wb.close()
    return salida
