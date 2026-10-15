"""Formatos en papel de la cátedra (UNIDAD 6) generados desde los datos del sistema [12.G].

- MANT/1CMP001 y MANT/2CMP001: calendario anual de preventivo (máquina × tarea × semana, con tiempo y lugar).
  El texto de las diapositivas no trae la grilla (es una imagen): se interpreta 1CMP001 = semanas 1 a 26 y
  2CMP001 = semanas 27 a 52 (B-017).
- MANT/MPV001: orden de trabajo de mantenimiento preventivo (UNIDAD 6, diap. 25).
- MANT/PMRED001: plantilla de mantenimiento predictivo, hasta 3 parámetros (UNIDAD 6, diap. 26).
"""

from __future__ import annotations

import io
from datetime import date, timedelta
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.orm import Session

from circularis_rel.datos.modelos import (
    Activo,
    OrdenTrabajo,
    OTTarea,
    PredictivoMedicion,
    PredictivoParametro,
    Usuario,
)
from circularis_rel.servicios.pdf import PDF
from circularis_rel.servicios.plan_preventivo import FilaCalendario, lunes

REFERENCIAS = {1: "MANT/1CMP001", 2: "MANT/2CMP001"}


def _semanas(semestre: int) -> range:
    return range(1, 27) if semestre == 1 else range(27, 53)


def pdf_calendario(filas: list[FilaCalendario], inicio: date) -> bytes:
    base = lunes(inicio)
    pdf = PDF("", orientacion="L")
    for semestre in (1, 2):
        pdf.titulo = "CALENDARIO ANUAL DE MANTENIMIENTO PREVENTIVO"
        pdf.referencia = REFERENCIAS[semestre]
        pdf.add_page()
        semanas = _semanas(semestre)
        desde = base + timedelta(weeks=semanas[0] - 1)
        hasta = base + timedelta(weeks=semanas[-1]) - timedelta(days=1)
        pdf.set_font("DejaVu", "", 8)
        pdf.cell(0, 5, f"Semanas {semanas[0]} a {semanas[-1]}: del {desde:%d/%m/%Y} al {hasta:%d/%m/%Y}",
                 new_x="LMARGIN", new_y="NEXT")
        anchos = [24, 28, 54, 25, 12] + [5] * 26  # 273 mm: A4 apaisado menos márgenes
        encabezado = ["Cód.", "Máquina", "Tarea", "Lugar", "Tiempo"] + [str(n) for n in semanas]
        cuerpo = [[f.maquina, f.nombre_maquina, f.tarea, f.lugar, "" if f.tiempo_h is None else f"{f.tiempo_h:g} h",
                   *["X" if n in f.semanas else "" for n in semanas]] for f in filas]
        pdf.tabla(encabezado, anchos, cuerpo, vacias=max(0, 8 - len(cuerpo)))
    return pdf.a_bytes()


def excel_calendario(filas: list[FilaCalendario], inicio: date) -> bytes:
    base = lunes(inicio)
    wb = Workbook()
    wb.remove(wb.active)
    borde = Border(*(Side(style="thin"),) * 4)
    for semestre in (1, 2):
        ws = wb.create_sheet(REFERENCIAS[semestre].replace("/", "-"))
        semanas = _semanas(semestre)
        ws.append([f"CALENDARIO ANUAL DE MANTENIMIENTO PREVENTIVO · Ref.: {REFERENCIAS[semestre]}"])
        ws.append([f"Semana 1 = {base:%d/%m/%Y}"])
        ws.append(["Cód. máquina", "Máquina", "Tarea", "Lugar", "Tiempo previsto (h)", *[f"S{n}" for n in semanas]])
        for f in filas:
            ws.append([f.maquina, f.nombre_maquina, f.tarea, f.lugar, f.tiempo_h, *["X" if n in f.semanas else None for n in semanas]])
        ws["A1"].font = Font(bold=True, size=12)
        for celda in ws[3]:
            celda.font = Font(bold=True)
            celda.fill = PatternFill("solid", fgColor="DDDDDD")
        for fila in ws.iter_rows(min_row=3, max_row=ws.max_row):
            for celda in fila:
                celda.border = borde
                if celda.column > 5:
                    celda.alignment = Alignment(horizontal="center")
        for i, ancho in enumerate([14, 26, 48, 18, 12], 1):
            ws.column_dimensions[get_column_letter(i)].width = ancho
        for i in range(6, 6 + len(semanas)):
            ws.column_dimensions[get_column_letter(i)].width = 4.5
    salida = io.BytesIO()
    wb.save(salida)
    return salida.getvalue()


def pdf_ot_preventiva(s: Session, id_ot: int, tz: ZoneInfo) -> bytes:
    ot = s.get(OrdenTrabajo, id_ot)
    activo = s.get(Activo, ot.id_activo)
    tareas = list(s.scalars(select(OTTarea).where(OTTarea.id_ot == ot.id).order_by(OTTarea.numero)))
    titulo = ("ORDEN DE TRABAJO PARA MANTENIMIENTO PREVENTIVO" if ot.tipo.value == "PREVENTIVA"
              else f"ORDEN DE TRABAJO · {ot.tipo.value.replace('_', ' ').capitalize()}")
    pdf = PDF(titulo, "MANT/MPV001")
    pdf.add_page()
    pdf.campo("N.º de OT", ot.numero)
    pdf.campo("Cód. máquina", activo.codigo)
    pdf.campo("Máquina", activo.nombre)
    pdf.campo("Operador", ot.responsable or "")
    pdf.campo("Tarea a realizar", tareas[0].descripcion if tareas else (ot.observaciones or ""))
    cuando = ot.inicio_real.astimezone(tz) if ot.inicio_real else None
    pdf.campo("Fecha", f"{cuando:%d/%m/%Y}" if cuando else (f"{ot.fecha_programada:%d/%m/%Y}" if ot.fecha_programada else ""))
    pdf.campo("Hora", f"{cuando:%H:%M}" if cuando else "")
    pdf.subtitulo("HERRAMIENTAS A UTILIZAR")
    pdf.set_font("DejaVu", "", 9)
    pdf.multi_cell(0, 6, ot.herramientas or " ", border=1, new_x="LMARGIN", new_y="NEXT")
    pdf.ln(2)
    pdf.tabla(["Descripción de la tarea", "Tarea N.º", "Observaciones", "Tiempo"], [92, 18, 58, 18],
              [[t.descripcion, t.numero, t.observaciones or "", "" if t.tiempo_h is None else f"{t.tiempo_h:g} h"] for t in tareas],
              vacias=max(2, 6 - len(tareas)))
    pdf.ln(10)
    pdf.set_font("DejaVu", "", 9)
    pdf.cell(90, 6, "Firma del operador: ____________________")
    pdf.cell(0, 6, "Controló: ____________________", new_x="LMARGIN", new_y="NEXT")
    return pdf.a_bytes()


def pdf_predictivo(s: Session, id_activo: int, especialidad: str | None, tz: ZoneInfo, filas_vacias: int = 10) -> bytes:
    activo = s.get(Activo, id_activo)
    consulta = select(PredictivoParametro).where(PredictivoParametro.id_activo == id_activo, PredictivoParametro.activo.is_(True))
    if especialidad:
        consulta = consulta.where(PredictivoParametro.especialidad == especialidad)
    parametros = list(s.scalars(consulta.order_by(PredictivoParametro.id)))[:3]
    pdf = PDF("PLANTILLA PARA MANTENIMIENTO PREDICTIVO", "MANT/PMRED001")
    pdf.add_page()
    pdf.campo("Especialidad", especialidad or "")
    pdf.campo("Código máq.", activo.codigo)
    pdf.campo("Máquina", activo.nombre)
    for i in range(3):
        p = parametros[i] if i < len(parametros) else None
        texto = (f"{p.nombre} ({p.unidad}); alerta {p.limite_alerta if p.limite_alerta is not None else '—'}, "
                 f"acción {p.limite_accion if p.limite_accion is not None else '—'}") if p else ""
        pdf.campo(f"Parámetro {i + 1}", texto)
    pdf.ln(3)
    filas = []
    por_fecha: dict = {}
    for i, p in enumerate(parametros):
        for m in s.scalars(select(PredictivoMedicion).where(PredictivoMedicion.id_parametro == p.id)):
            clave = (m.fecha, m.equipo_medicion, m.id_usuario_controlo)
            por_fecha.setdefault(clave, [None, None, None])[i] = m.valor
    for (fecha, equipo, controlo), valores in sorted(por_fecha.items(), key=lambda kv: kv[0][0]):
        nombre = s.get(Usuario, controlo).nombre if controlo else ""
        filas.append([f"{fecha.astimezone(tz):%d/%m/%Y}", equipo or "", *["" if v is None else f"{v:g}" for v in valores], nombre])
    pdf.tabla(["Fecha", "Eq. medic.", "Med. parám. 1", "Med. parám. 2", "Med. parám. 3", "Controló"], [24, 34, 30, 30, 30, 38],
              filas, vacias=filas_vacias)
    return pdf.a_bytes()
